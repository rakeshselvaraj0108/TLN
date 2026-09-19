"""Read a case document in any common format and report what it names.

Reading is separate from ingesting: analysing a complaint does not enter it into evidence. When a
document contains a statement-shaped table its rows are returned in the bank schema so the caller
can choose to ingest them through the same path a CSV takes.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from xml.etree import ElementTree

from tracex_api import db
from tracex_api.engine import scoring
from tracex_api.engine.dataset import Dataset

PATTERNS = [
    ("phone", re.compile(r"(?<![\w+])(?:\+91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)")),
    ("imei", re.compile(r"(?<!\d)\d{15}(?!\d)")),
    ("ifsc", re.compile(r"\b[A-Z]{4}0[A-Z0-9]{6}\b")),
    ("account", re.compile(r"\b[A-Z]{4}\d{11}\b")),
    ("upi", re.compile(r"\b[\w.\-]{2,}@(?:ok\w+|ybl|paytm|upi|ibl|axl|apl|icici|hdfcbank|sbi|kotak|yesbank)\b", re.I)),
    ("email", re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.]{2,}\b")),
    ("ip", re.compile(r"\b(?:(?:25[0-5]|2[0-4]\d|1?\d?\d)\.){3}(?:25[0-5]|2[0-4]\d|1?\d?\d)\b")),
    ("vehicle", re.compile(r"\b[A-Z]{2}\s?\d{2}\s?[A-Z]{1,2}\s?\d{4}\b")),
    ("case_code", re.compile(r"\bTRX-\d{4}-\d{4}\b")),
    ("record_id", re.compile(r"\b(?:CDR|IPDR|TXN|SOC|ALP)\d{5}\b")),
    ("handle", re.compile(r"(?<![\w@])@[A-Za-z_][\w]{2,30}\b")),
]
AMOUNT = re.compile(r"(?:₹|Rs\.?|INR)\s?\d[\d,]*(?:\.\d{1,2})?(?:\s?(?:lakh|lac|crore|cr|k))?|\b\d+(?:\.\d+)?\s?(?:lakh|lac|crore)\b", re.I)
DATE = re.compile(r"\b(?:\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s\d{4})\b")
FLAGS = [
    ("digital arrest", r"digital\s+arrest"), ("impersonation", r"\b(?:CBI|ED|customs|narcotics|police)\s+officer|fake\s+(?:CBI|police)"),
    ("money mule", r"account\s+rent|rent\s+(?:your|bank)\s+account|mule"), ("urgency", r"\burgent(?:ly)?\b|immediately|within\s+\d+\s+(?:hours|minutes)"),
    ("remote access", r"anydesk|teamviewer|screen\s*share|quick\s*support"), ("OTP request", r"\bOTP\b"),
    ("security deposit", r"security\s+deposit|clearance\s+fee|verification\s+fee"), ("crypto", r"\b(?:USDT|crypto|bitcoin|wallet\s+address)\b"),
    ("SIM swap", r"sim\s+swap|new\s+sim|e-?sim"),
]
STATEMENT_ALIASES = {
    "txn_time": ["txn_time", "date", "txn date", "transaction date", "value date", "posting date", "datetime", "time"],
    "amount": ["amount", "amt", "transaction amount"],
    "debit": ["debit", "withdrawal", "withdrawal amt", "dr"],
    "credit": ["credit", "deposit", "deposit amt", "cr"],
    "narration": ["narration", "description", "particulars", "remarks", "details"],
    "balance_after": ["balance", "balance_after", "closing balance", "running balance"],
    "channel": ["channel", "mode", "type"],
    "record_id": ["record_id", "ref", "reference", "ref no", "chq/ref no", "utr", "txn id", "transaction id"],
    "src_account": ["src_account", "from account", "account", "account number"],
    "dst_account": ["dst_account", "to account", "beneficiary account", "counterparty account"],
    "src_name": ["src_name", "account holder", "name"],
    "dst_name": ["dst_name", "beneficiary", "beneficiary name", "counterparty"],
}
BANK_COLUMNS = ["record_id", "txn_time", "amount", "direction", "channel", "src_account", "src_ifsc", "src_name", "dst_account", "dst_ifsc", "dst_name", "narration", "balance_after"]


def _extract(content: bytes, filename: str) -> tuple[str, str, str, int, list[list[list[str]]], list[str], bool]:
    """-> text, method, note, page_count, tables, warnings, is_transcription"""
    name = filename.lower()
    warnings: list[str] = []
    if name.endswith(".pdf") or content[:5] == b"%PDF-":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = [p.extract_text() or "" for p in reader.pages]
            text = "\n".join(pages).strip()
            tables = [_table_from_lines(p.splitlines()) for p in pages]
            tables = [t for t in tables if t]
            if not text:
                return "", "unreadable", "scanned PDF with no text layer — OCR is not installed on this deployment", len(pages), [], ["no text layer"], False
            if tables:
                return text, "text_layer", "read exactly from the PDF's text layer", len(pages), tables, warnings, False
            return text, "narrative", "read as prose — no ruled table in this document", len(pages), [], warnings, False
        except Exception as exc:  # never fail a read
            return "", "unreadable", f"could not parse PDF: {exc}", 0, [], [str(exc)], False
    if name.endswith(".docx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                xml = ElementTree.fromstring(z.read("word/document.xml"))
            ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            paragraphs = ["".join(t.text or "" for t in p.iter(f"{ns}t")) for p in xml.iter(f"{ns}p")]
            tables = []
            for tbl in xml.iter(f"{ns}tbl"):
                rows = [["".join(t.text or "" for t in c.iter(f"{ns}t")).strip() for c in r.iter(f"{ns}tc")] for r in tbl.iter(f"{ns}tr")]
                if len(rows) > 1:
                    tables.append(rows)
            return "\n".join(paragraphs).strip(), "docx", "read from the Word document's text", 0, tables, warnings, False
        except Exception as exc:
            return "", "unreadable", f"could not parse Word document: {exc}", 0, [], [str(exc)], False
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            tables, lines = [], []
            for ws in wb.worksheets:
                rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
                rows = [r for r in rows if any(c.strip() for c in r)]
                if rows:
                    tables.append(rows)
                    lines += [" ".join(r) for r in rows]
            return "\n".join(lines), "spreadsheet", "read cell by cell from the workbook", len(wb.worksheets), tables, warnings, False
        except Exception as exc:
            return "", "unreadable", f"could not parse workbook: {exc}", 0, [], [str(exc)], False
    if name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
        return "", "unreadable", "image — OCR is not installed on this deployment, so no text could be transcribed", 1, [], ["OCR unavailable"], True
    text = content.decode("utf-8-sig", errors="replace")
    if name.endswith((".json", ".ndjson", ".jsonl")):
        try:
            data = [json.loads(l) for l in text.splitlines() if l.strip()] if name.endswith((".ndjson", ".jsonl")) else json.loads(text)
            items = data if isinstance(data, list) else [data]
            dicts = [i for i in items if isinstance(i, dict)]
            if dicts:
                cols = list(dict.fromkeys(k for d in dicts for k in d))
                return text, "structured", "read as structured JSON records", 0, [[cols] + [[str(d.get(c, "")) for c in cols] for d in dicts]], warnings, False
        except ValueError as exc:
            warnings.append(f"not valid JSON: {exc}")
    if name.endswith((".csv", ".tsv")) or ("," in text.splitlines()[0] if text.strip() else False):
        dialect = "excel-tab" if name.endswith(".tsv") else "excel"
        rows = [r for r in csv.reader(io.StringIO(text), dialect=dialect) if any(c.strip() for c in r)]
        if len(rows) > 1:
            return text, "delimited", "read as delimited rows", 0, [rows], warnings, False
    return text, "plain_text", "read as plain text", 0, [], warnings, False


def _table_from_lines(lines: list[str]) -> list[list[str]] | None:
    rows = [re.split(r"\s{2,}|\t|\s\|\s", l.strip()) for l in lines if l.strip()]
    rows = [r for r in rows if len(r) >= 3]
    return rows if len(rows) >= 2 else None


def _statement_rows(tables: list[list[list[str]]], warnings: list[str]) -> tuple[list[dict], list[str]]:
    out: list[dict] = []
    for table in tables:
        header = [h.strip().lower() for h in table[0]]
        mapping = {}
        for field, aliases in STATEMENT_ALIASES.items():
            for i, h in enumerate(header):
                if h in aliases and field not in mapping:
                    mapping[field] = i
        if "txn_time" not in mapping or not ({"amount", "debit", "credit"} & mapping.keys()):
            warnings.append(f"no recognised statement columns in {table[0][:3]}")
            continue
        for n, raw in enumerate(table[1:], start=1):
            get = lambda f: raw[mapping[f]].strip() if f in mapping and mapping[f] < len(raw) else ""
            amount, direction = get("amount"), "DEBIT"
            if not amount:
                debit, credit = get("debit"), get("credit")
                amount, direction = (debit, "DEBIT") if debit else (credit, "CREDIT")
            amount = re.sub(r"[^\d.]", "", amount)
            if not amount:
                continue
            row = {c: "" for c in BANK_COLUMNS}
            row.update({f: get(f) for f in mapping if f in row})
            row["amount"], row["direction"] = amount, direction
            row["record_id"] = row["record_id"] or f"DOC{n:05d}"
            out.append(row)
    return out, (BANK_COLUMNS if out else [])


def _context(text: str, start: int, end: int, width: int = 60) -> str:
    return re.sub(r"\s+", " ", text[max(0, start - width):min(len(text), end + width)]).strip()


def analyze(content: bytes, filename: str, ds: Dataset) -> dict:
    text, method, note, pages, tables, warnings, transcription = _extract(content, filename)
    scores = scoring.scores(ds)
    known_values = {
        "phone": ds.by_phone, "imei": ds.by_device, "account": ds.by_account, "handle": ds.by_handle, "vehicle": ds.by_plate,
    }
    found: dict[tuple[str, str], dict] = {}
    claimed: list[tuple[int, int]] = []
    for kind, pattern in PATTERNS:
        for m in pattern.finditer(text):
            if any(s <= m.start() < e for s, e in claimed):
                continue
            value = re.sub(r"[\s-]", "", m.group(0)) if kind in ("phone", "vehicle") else m.group(0)
            if kind == "phone" and not value.startswith("+91"):
                value = "+91" + value[-10:]
            key = (kind, value)
            if key in found:
                found[key]["count"] += 1
                continue
            claimed.append((m.start(), m.end()))
            entity_id = known_values.get(kind, {}).get(value)
            known = entity_id is not None or (kind == "record_id" and any(value in idx for idx in (ds.calls_by_id, ds.txns_by_id, ds.sessions_by_id, ds.posts_by_id, ds.reads_by_id)))
            if kind == "ip":
                known = known or any(s.public_ip == value or s.dest_ip == value for s in ds.sessions)
            s = scores.get(entity_id) if entity_id else None
            found[key] = {
                "kind": kind, "value": value, "count": 1, "context": _context(text, m.start(), m.end()), "known": bool(known),
                "entity_id": entity_id, "band": s["band"] if s else None, "risk_score": round(s["risk_score"] * 100, 1) if s else None,
                "model": s["model"] if s else None,
            }
    hits = list(found.values())
    amounts = [{"text": m.group(0), "context": _context(text, m.start(), m.end(), 40)} for m in AMOUNT.finditer(text)][:50]
    dates = list(dict.fromkeys(m.group(0) for m in DATE.finditer(text)))[:50]
    flags = []
    for label, pattern in FLAGS:
        m = re.search(pattern, text, re.I)
        if m:
            flags.append({"label": label, "context": _context(text, m.start(), m.end(), 80)})
    rows, row_columns = _statement_rows(tables, warnings)
    resolved = sorted({h["entity_id"] for h in hits if h["entity_id"]})
    known_scores = [h["risk_score"] for h in hits if h["risk_score"] is not None]
    return {
        "filename": filename,
        "size_bytes": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
        "extraction_method": method,
        "extraction_note": note,
        "is_transcription": transcription,
        "page_count": pages,
        "mean_ocr_confidence": None,
        "char_count": len(text),
        "text_preview": text[:4000],
        "hits": hits,
        "amounts": amounts,
        "dates": dates,
        "flags": flags,
        "rows": rows[:500],
        "row_columns": row_columns,
        "warnings": warnings[:20],
        "analysed_at": datetime.now(timezone.utc).isoformat(),
        "summary": {
            "identifiers": len(hits),
            "by_kind": dict(Counter(h["kind"] for h in hits)),
            "already_known": sum(1 for h in hits if h["known"]),
            "resolved_entities": resolved,
            "highest_known_score": max(known_scores) if known_scores else None,
            "models": sorted({h["model"] for h in hits if h["model"]}),
            "rows_parsed": len(rows),
        },
        "identifier_count": len(hits),
        "ingested": None,
    }


def store(analysis: dict, user: str, case_id: int | None = None, ingested_batch_id: str | None = None) -> dict:
    doc_id = db.next_id("documents")
    entities = analysis["summary"]["resolved_entities"]
    doc = {
        "id": doc_id,
        "filename": analysis["filename"],
        "sha256": analysis["sha256"],
        "size_bytes": analysis["size_bytes"],
        "extraction_method": analysis["extraction_method"],
        "page_count": analysis["page_count"],
        "char_count": analysis["char_count"],
        "identifier_count": analysis["summary"]["identifiers"],
        "known_count": analysis["summary"]["already_known"],
        "entities": entities,
        "top_risk": analysis["summary"]["highest_known_score"] or 0.0,
        "case_id": case_id,
        "ingested_batch_id": ingested_batch_id,
        "uploaded_by": user,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "analysis": {k: v for k, v in analysis.items() if k != "identifier_count"},
    }
    db.doc_put("documents", doc_id, doc)
    return doc
